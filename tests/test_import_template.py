"""Tests for Excel fleet import template and parser consistency."""

import io
import os
from fastapi.testclient import TestClient
import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.auth.providers import LocalAuthProvider
from app.core.database import Base, get_db
from app.importer.schema import FLEET_IMPORT_COLUMNS
from app.importer.template import TEMPLATE_FILENAME, build_template
from app.main import app
from app.models.models import User, UserRole
from app.routers.import_fleet import parse_excel_rows

SQLALCHEMY_DATABASE_URL = "sqlite:///./test_import_template.db"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False}
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

@pytest.fixture(scope="module")
def setup_db():
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)
    if os.path.exists("./test_import_template.db"):
        os.remove("./test_import_template.db")

@pytest.fixture
def db_session(setup_db):
    connection = engine.connect()
    transaction = connection.begin()
    session = TestingSessionLocal(bind=connection)
    yield session
    session.close()
    transaction.rollback()
    connection.close()


def test_schema_constant_structure():
    """Verify schema constant FLEET_IMPORT_COLUMNS contains required fields."""
    assert len(FLEET_IMPORT_COLUMNS) >= 8
    keys = [col.key for col in FLEET_IMPORT_COLUMNS]
    assert "hostname" in keys
    assert "ip" in keys
    assert "mac" in keys
    assert "is_round_the_clock" in keys

    hostname_spec = next(col for col in FLEET_IMPORT_COLUMNS if col.key == "hostname")
    assert hostname_spec.required is True
    assert hostname_spec.header_en == "hostname"


def test_template_generator_output():
    """Verify build_template produces valid openpyxl readable workbook."""
    content = build_template()
    assert isinstance(content, bytes)
    assert len(content) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "Computers" in wb.sheetnames
    assert "Readme" in wb.sheetnames

    sheet_comp = wb["Computers"]
    headers = [cell.value for cell in sheet_comp[1]]
    expected_headers = [col.header_en for col in FLEET_IMPORT_COLUMNS]
    assert headers == expected_headers

    example_row = [cell.value for cell in sheet_comp[2]]
    assert example_row[0] == "PC-OFFICE-101"


def test_template_roundtrip_parsing():
    """Verify that generated template example row parses cleanly with 0 errors."""
    content = build_template()
    parsed_rows = parse_excel_rows(content)

    assert len(parsed_rows) == 1
    row = parsed_rows[0]
    assert row["hostname"] == "PC-OFFICE-101"
    assert row["ip"] == "192.168.1.50"
    assert row["mac"] == "00:11:22:33:44:55"


def test_download_template_access_control(db_session):
    """Verify route GET /admin/import/template returns 200 for ADMIN, 403 for non-admins."""
    provider = LocalAuthProvider()

    admin = User(
        username="admin_import_test",
        email_or_login="admin_import_test@cfms.local",
        password_hash=provider.hash_password("adminpass"),
        role=UserRole.ADMIN,
        is_active=True,
    )
    tech = User(
        username="tech_import_test",
        email_or_login="tech_import_test@cfms.local",
        password_hash=provider.hash_password("techpass"),
        role=UserRole.TECHNICIAN,
        is_active=True,
    )
    user = User(
        username="user_import_test",
        email_or_login="user_import_test@cfms.local",
        password_hash=provider.hash_password("userpass"),
        role=UserRole.USER,
        is_active=True,
    )
    db_session.add_all([admin, tech, user])
    db_session.commit()

    def _override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = _override_get_db
    client = TestClient(app)

    # ADMIN
    client.post("/auth/login", data={"username": "admin_import_test", "password": "adminpass"})
    res_admin = client.get("/admin/import/template")
    assert res_admin.status_code == 200
    assert res_admin.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f'filename="{TEMPLATE_FILENAME}"' in res_admin.headers["content-disposition"]
    assert len(res_admin.content) > 1000

    # TECHNICIAN -> 403
    client.post("/auth/login", data={"username": "tech_import_test", "password": "techpass"})
    res_tech = client.get("/admin/import/template")
    assert res_tech.status_code == 403

    # USER -> 403
    client.post("/auth/login", data={"username": "user_import_test", "password": "userpass"})
    res_user = client.get("/admin/import/template")
    assert res_user.status_code == 403

    app.dependency_overrides.clear()


def test_preview_and_confirm_generated_template(db_session):
    """Test uploading generated template via UI preview and confirming import."""
    provider = LocalAuthProvider()
    admin = User(
        username="admin",
        email_or_login="admin@cfms.local",
        password_hash=provider.hash_password("adminpass"),
        role=UserRole.ADMIN,
        is_active=True,
    )
    db_session.add(admin)
    db_session.commit()

    def _override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = _override_get_db
    client = TestClient(app)
    client.post("/auth/login", data={"username": "admin", "password": "adminpass"})

    template_bytes = build_template()

    # Post template to /admin/import/preview
    files = {"file": (TEMPLATE_FILENAME, template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res_preview = client.post("/admin/import/preview", files=files)
    assert res_preview.status_code == 200
    assert "PC-OFFICE-101" in res_preview.text
    assert "Готов к импорту" in res_preview.text
    assert "file_token" in res_preview.text

    # Extract file_token from HTML hidden input
    import re
    match = re.search(r'name="file_token"\s+value="([^"]+)"', res_preview.text)
    assert match is not None
    file_token = match.group(1)

    # Confirm import
    res_confirm = client.post("/admin/import/confirm", data={"file_token": file_token}, follow_redirects=True)
    assert res_confirm.status_code == 200
    assert "Импорт завершен" in res_confirm.text

    app.dependency_overrides.clear()
