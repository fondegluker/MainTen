"""Generator for Excel fleet import template."""

import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.import.schema import FLEET_IMPORT_COLUMNS

TEMPLATE_FILENAME = "fleet_import_template.xlsx"
CANONICAL_ASSET_PATH = Path(__file__).parent / "assets" / TEMPLATE_FILENAME


def build_template() -> bytes:
    wb = openpyxl.Workbook()

    # Sheet 1: Computers
    ws_computers = wb.active
    ws_computers.title = "Computers"

    # Sheet 2: Readme
    ws_readme = wb.create_sheet(title="Readme")

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")

    # Populate Sheet 1
    headers = [col.header_en for col in FLEET_IMPORT_COLUMNS]
    example_row = [col.example_value for col in FLEET_IMPORT_COLUMNS]

    ws_computers.append(headers)
    ws_computers.append(example_row)

    # Style header row
    for col_num, _ in enumerate(headers, 1):
        cell = ws_computers.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Freeze header row
    ws_computers.freeze_panes = "A2"

    # Set column widths
    for col_num, col_spec in enumerate(FLEET_IMPORT_COLUMNS, 1):
        col_letter = get_column_letter(col_num)
        max_len = max(len(col_spec.header_en), len(str(col_spec.example_value)), 12)
        ws_computers.column_dimensions[col_letter].width = max_len + 5

    # Populate Sheet 2 (Readme / Legend)
    readme_headers = [
        "Column / Колонка",
        "Target Field / Поле БД",
        "Type / Тип",
        "Required / Обязательное",
        "Example / Пример",
        "Description (RU)",
        "Description (EN)",
    ]
    ws_readme.append(readme_headers)

    for col_num in range(1, len(readme_headers) + 1):
        cell = ws_readme.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in FLEET_IMPORT_COLUMNS:
        ws_readme.append([
            col.header_en,
            col.field_name,
            col.type_name,
            "Да / Yes" if col.required else "Нет / No",
            col.example_value,
            col.description_ru,
            col.description_en,
        ])

    ws_readme.freeze_panes = "A2"

    for col_num in range(1, len(readme_headers) + 1):
        col_letter = get_column_letter(col_num)
        ws_readme.column_dimensions[col_letter].width = 24

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()

    # Ensure canonical asset directory exists and file is updated/saved
    CANONICAL_ASSET_PATH.parent.mkdir(parents=True, exist_ok=True)
    CANONICAL_ASSET_PATH.write_bytes(content)

    return content


if __name__ == "__main__":
    build_template()
