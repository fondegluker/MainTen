import io
import uuid
from typing import Any

import openpyxl
from fastapi import APIRouter, Depends, File, Form, Request, UploadFile, status
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.auth.dependencies import require_admin
from app.core.database import get_db
from app.core.validators import validate_ip, validate_mac
from app.models.models import Computer, User
from app.routers.web import context_with_defaults
from app.services.audit_service import log_audit

router = APIRouter(prefix="/admin/import", tags=["admin-import"])
templates = Jinja2Templates(directory="app/templates")

IMPORT_STAGING_CACHE: dict[str, dict[str, Any]] = {}

def parse_excel_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    parsed = []

    for idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        row_dict = {}
        for col_name, val in zip(header, row, strict=False):
            row_dict[col_name] = str(val).strip() if val is not None else ""
        row_dict["_row_idx"] = idx
        parsed.append(row_dict)
    return parsed

@router.get("", response_class=HTMLResponse)
def import_page(
    request: Request,
    error: str | None = None,
    current_user: User = Depends(require_admin)
):
    return templates.TemplateResponse(
        "admin/import.html",
        context_with_defaults(request, current_user, {"preview_rows": None, "error": error})
    )

@router.get("/template")
def download_import_template(
    current_user: User = Depends(require_admin)
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fleet Import Template"

    headers = ["hostname", "ip", "mac", "os", "location", "owner", "is_round_the_clock", "notes"]
    ws.append(headers)

    # Example row
    example_row = ["PC-OFFICE-101", "192.168.1.50", "00:11:22:33:44:55", "Windows 11", "Room 302", "admin", "нет", "Рабочий ПК"]
    ws.append(example_row)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers_resp = {
        'Content-Disposition': 'attachment; filename="fleet_import_template.xlsx"'
    }
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_resp
    )

@router.post("/preview", response_class=HTMLResponse)
async def preview_import(
    request: Request,
    file: UploadFile = File(...),
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith(".xlsx"):
        return templates.TemplateResponse(
            "admin/import.html",
            context_with_defaults(request, current_user, {"preview_rows": None, "error": "Пожалуйста, загрузите файл формата .xlsx"}),
            status_code=status.HTTP_400_BAD_REQUEST
        )

    content = await file.read()
    raw_rows = parse_excel_rows(content)

    existing_computers = {c.hostname.lower(): c for c in db.query(Computer).all()}
    existing_users = {u.username.lower(): u for u in db.query(User).all()}
    existing_users.update({u.email_or_login.lower(): u for u in db.query(User).all()})

    preview_rows = []
    valid_rows_for_import = []
    seen_hostnames_in_file = set()
    has_errors = False

    for r in raw_rows:
        hostname = r.get("hostname") or r.get("компьютер") or r.get("имя хоста") or ""
        ip = r.get("ip") or r.get("ip-адрес") or ""
        mac = r.get("mac") or r.get("mac-адрес") or ""
        os_name = r.get("os") or r.get("операционная система") or ""
        location = r.get("location") or r.get("кабинет") or r.get("расположение") or ""
        owner_text = r.get("owner") or r.get("владелец") or r.get("пользователь") or ""
        notes = r.get("notes") or r.get("заметки") or ""
        rtc_val = str(r.get("is_round_the_clock") or r.get("24/7") or "").lower()

        is_rtc = rtc_val in ("1", "true", "да", "yes")

        errors = []
        warning = None
        diff_type = "new"

        if not hostname:
            errors.append("Отсутствует hostname")
        elif hostname.lower() in seen_hostnames_in_file:
            errors.append(f"Дубликат hostname '{hostname}' в файле")
        else:
            seen_hostnames_in_file.add(hostname.lower())

        if ip and not validate_ip(ip):
            errors.append(f"Некорректный IP '{ip}'")

        if mac and not validate_mac(mac):
            errors.append(f"Некорректный MAC '{mac}'")

        # Owner lookup
        owner_user_id = None
        if owner_text:
            matched_user = existing_users.get(owner_text.lower())
            if matched_user:
                owner_user_id = matched_user.id
            else:
                warning = f"Пользователь '{owner_text}' не найден (будет не назначен)"

        # Diff detection
        if hostname and hostname.lower() in existing_computers:
            diff_type = "update"
            if not warning:
                warning = f"Компьютер '{hostname}' уже существует (будет обновлен)"

        is_valid = len(errors) == 0
        if not is_valid:
            has_errors = True

        parsed_row = {
            "row_idx": r["_row_idx"],
            "hostname": hostname,
            "ip": ip,
            "mac": mac,
            "os": os_name,
            "location": location,
            "owner_text": owner_text,
            "owner_user_id": owner_user_id,
            "is_round_the_clock": is_rtc,
            "notes": notes,
            "diff_type": diff_type,
            "is_valid": is_valid,
            "errors": errors,
            "warning": warning,
        }
        preview_rows.append(parsed_row)
        if is_valid:
            valid_rows_for_import.append(parsed_row)

    file_token = str(uuid.uuid4())
    IMPORT_STAGING_CACHE[file_token] = {
        "raw_filename": file.filename,
        "rows": valid_rows_for_import,
        "has_errors": has_errors
    }

    return templates.TemplateResponse(
        "admin/import.html",
        context_with_defaults(request, current_user, {
            "preview_rows": preview_rows,
            "file_token": file_token,
            "has_errors": has_errors,
            "error": None
        })
    )

@router.post("/confirm")
def confirm_import(
    file_token: str = Form(...),
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    staging_data = IMPORT_STAGING_CACHE.pop(file_token, None)
    if not staging_data:
        return RedirectResponse(url="/admin/import?error=Сессия+импорта+истекла", status_code=status.HTTP_302_FOUND)

    if staging_data.get("has_errors"):
        return RedirectResponse(url="/admin/import?error=Файл+содержит+критические+ошибки+и+отклонен", status_code=status.HTTP_302_FOUND)

    staged_rows = staging_data["rows"]
    raw_filename = staging_data["raw_filename"]

    imported_count = 0
    updated_count = 0

    try:
        with db.begin_nested():
            for r in staged_rows:
                hostname = r["hostname"]
                comp = db.query(Computer).filter(func.lower(Computer.hostname) == hostname.lower()).first()
                if comp:
                    comp.ip = r["ip"] or comp.ip
                    comp.mac = r["mac"] or comp.mac
                    comp.os = r["os"] or comp.os
                    comp.location = r["location"] or comp.location
                    comp.owner_user_id = r["owner_user_id"] or comp.owner_user_id
                    comp.is_round_the_clock = r["is_round_the_clock"]
                    comp.notes = r["notes"] or comp.notes
                    updated_count += 1
                else:
                    comp = Computer(
                        hostname=hostname,
                        ip=r["ip"] or None,
                        mac=r["mac"] or None,
                        os=r["os"] or None,
                        location=r["location"] or None,
                        owner_user_id=r["owner_user_id"],
                        is_round_the_clock=r["is_round_the_clock"],
                        notes=r["notes"] or None,
                        status="active"
                    )
                    db.add(comp)
                    imported_count += 1
        db.commit()
    except Exception as exc:
        db.rollback()
        return RedirectResponse(url=f"/admin/import?error=Ошибка+транзакции+импорта:+{exc}", status_code=status.HTTP_302_FOUND)

    log_audit(
        db=db,
        actor_user_id=current_user.id,
        action="excel_import_fleet",
        entity="computers",
        after={
            "raw_filename": raw_filename,
            "imported_count": imported_count,
            "updated_count": updated_count,
            "total_processed": len(staged_rows)
        }
    )

    return RedirectResponse(
        url=f"/admin/computers?message=Импорт+завершен:+создано+{imported_count},+обновлено+{updated_count}",
        status_code=status.HTTP_302_FOUND
    )
